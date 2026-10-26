import tempfile
import unittest
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.config import load_settings
from app.extractors import ExtractionError, extract
from app.ingest import FileMeta, ingest_file
from app.rag import RagService
from app.store import MemoryStore
from tests.helpers import FakeEmbedder, FakeLLM


def make_pptx(path):
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Symmetric Encryption"
    s1.placeholders[1].text = "Symmetric encryption uses a single shared secret key for both encryption and decryption."
    s1.notes_slide.notes_text_frame.text = "Mention AES and DES as examples."
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Asymmetric Encryption"
    box = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    box.text_frame.text = "Asymmetric encryption uses a public key and a private key pair such as RSA."
    prs.save(path)


def make_docx(path):
    d = Document()
    d.add_heading("Firewalls", level=1)
    d.add_paragraph("A firewall filters traffic between trusted and untrusted networks using rules.")
    d.add_heading("Intrusion Detection", level=1)
    d.add_paragraph("An IDS monitors network traffic for suspicious activity and raises alerts.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Type"
    t.rows[0].cells[1].text = "Signature based"
    d.save(path)


class IngestTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_pptx_extraction_has_titles_notes_and_slide_numbers(self):
        p = self.dir / "a.pptx"
        make_pptx(p)
        segs = extract(p, "pptx")
        self.assertEqual([s.number for s in segs], [1, 2])
        self.assertEqual(segs[0].label, "Symmetric Encryption")
        self.assertIn("single shared secret key", segs[0].text)
        self.assertIn("Speaker notes: Mention AES", segs[0].text)
        self.assertIn("RSA", segs[1].text)

    def test_docx_extraction_splits_on_headings_and_includes_tables(self):
        p = self.dir / "a.docx"
        make_docx(p)
        segs = extract(p, "docx")
        self.assertEqual([s.label for s in segs], ["Firewalls", "Intrusion Detection"])
        self.assertIn("Type | Signature based", segs[1].text)

    def test_unsupported_and_corrupt_files_give_clear_errors(self):
        bad = self.dir / "bad.docx"
        bad.write_bytes(b"not a real docx")
        with self.assertRaises(ExtractionError):
            extract(bad, "docx")
        with self.assertRaises(ExtractionError):
            extract(bad, "exe")

    def test_image_without_ocr_is_rejected_clearly(self):
        from PIL import Image
        img = self.dir / "x.png"
        Image.new("RGB", (20, 20), "white").save(img)
        with self.assertRaises(ExtractionError) as cm:
            extract(img, "png", ocr=False)
        self.assertIn("OCR", str(cm.exception))

    def test_full_pipeline_ingest_then_ask_with_citation(self):
        p = self.dir / "a.pptx"
        make_pptx(p)
        settings = load_settings({"AI_SERVICE_KEY": "k"})
        emb, store = FakeEmbedder(), MemoryStore()
        meta = FileMeta(7, 3, "IAS2", 4, 1, "Encryption Basics", "a.pptx")
        out = ingest_file(p, "pptx", meta, settings, emb, store)
        self.assertEqual(out["file_id"], 7)
        self.assertGreaterEqual(out["chunks"], 1)

        rag = RagService(emb, store, FakeLLM("RSA uses a public and private key pair [1]."), 0.2, 5)
        ans = rag.ask("What keys does asymmetric encryption RSA use?", [3])
        self.assertTrue(ans.found)
        self.assertEqual(ans.sources[0].file_name, "a.pptx")
        self.assertTrue(ans.sources[0].locator.startswith("Slide"))

        # re-ingesting replaces, never duplicates; deleting removes everything
        ingest_file(p, "pptx", meta, settings, emb, store)
        self.assertEqual(store.count(), out["chunks"])
        store.delete_file(7)
        self.assertEqual(store.count(), 0)
        self.assertFalse(rag.ask("asymmetric encryption RSA key", [3]).found)


if __name__ == "__main__":
    unittest.main()
