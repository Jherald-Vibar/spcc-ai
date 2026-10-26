"""Turn an uploaded file into text segments (one per PDF page / PPTX slide / DOCX section / image)."""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

SUPPORTED = {"pdf", "pptx", "docx", "png", "jpg", "jpeg"}
MIN_TEXT_CHARS = 25  # a PDF page with less text than this is treated as scanned


class ExtractionError(Exception):
    """The file could not be read. The message is shown to the teacher."""


@dataclass
class Segment:
    unit: str      # page | slide | section | image
    number: int    # 1-based position inside the file
    label: str     # slide title / section heading, or ""
    text: str


def clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract(path: Path, extension: str, *, ocr: bool = True, ocr_lang: str = "eng") -> list[Segment]:
    ext = extension.lower().lstrip(".")
    if ext not in SUPPORTED:
        raise ExtractionError(f"Unsupported file type: .{ext}")
    if ext == "pdf":
        segments = _pdf(path, ocr, ocr_lang)
    elif ext == "pptx":
        segments = _pptx(path)
    elif ext == "docx":
        segments = _docx(path)
    else:
        segments = _image(path, ocr, ocr_lang)

    if not segments:
        hint = "" if ocr else " Scanned files need OCR (OCR_ENABLED=true and Tesseract installed)."
        raise ExtractionError("No readable text was found in this file." + hint)
    return segments


# ---------- OCR helpers ----------

def _ocr_image(image, lang: str) -> str:
    try:
        import pytesseract
    except ImportError as exc:
        raise ExtractionError("OCR needs the pytesseract package.") from exc
    try:
        return pytesseract.image_to_string(image, lang=lang)
    except pytesseract.TesseractNotFoundError as exc:
        raise ExtractionError(
            "OCR needs the Tesseract program installed on the AI service machine."
        ) from exc
    except pytesseract.TesseractError as exc:
        raise ExtractionError(f"OCR failed (is the '{lang}' language data installed?).") from exc


# ---------- PDF ----------

def _pdf(path: Path, ocr: bool, ocr_lang: str) -> list[Segment]:
    import fitz  # PyMuPDF

    try:
        doc = fitz.open(path)
    except Exception as exc:
        raise ExtractionError("The PDF could not be opened (it may be corrupted).") from exc

    segments: list[Segment] = []
    with doc:
        if doc.needs_pass:
            raise ExtractionError("The PDF is password protected.")
        for index, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if len(text) < MIN_TEXT_CHARS and ocr:
                from PIL import Image

                pix = page.get_pixmap(dpi=200)
                image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                text = _ocr_image(image, ocr_lang).strip() or text
            text = clean(text)
            if text:
                segments.append(Segment("page", index, "", text))
    return segments


# ---------- PPTX ----------

def _shape_text(shape) -> list[str]:
    parts: list[str] = []
    if getattr(shape, "shape_type", None) is not None and shape.shape_type == 6:  # GROUP
        for child in shape.shapes:
            parts.extend(_shape_text(child))
        return parts
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            parts.append(text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return parts


def _pptx(path: Path) -> list[Segment]:
    from pptx import Presentation

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        raise ExtractionError("The PowerPoint file could not be opened.") from exc

    segments: list[Segment] = []
    for index, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            lines.extend(_shape_text(shape))

        title = ""
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text_frame.text.strip()
        except Exception:
            title = ""

        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = frame.text.strip() if frame is not None else ""
            if notes:
                lines.append("Speaker notes: " + notes)

        text = clean("\n".join(lines))
        if text:
            segments.append(Segment("slide", index, title, text))
    return segments


# ---------- DOCX ----------

def _docx(path: Path) -> list[Segment]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(str(path))
    except Exception as exc:
        raise ExtractionError("The Word file could not be opened.") from exc

    sections: list[tuple[str, list[str]]] = []
    label, lines = "", []

    def flush() -> None:
        nonlocal lines
        if lines:
            sections.append((label, lines))
        lines = []

    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            style = (block.style.name if block.style is not None else "") or ""
            if style.startswith("Heading") or style == "Title":
                flush()
                label = text
            lines.append(text)
        elif isinstance(block, Table):
            for row in block.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
    flush()

    segments = []
    for number, (heading, body) in enumerate(sections, start=1):
        text = clean("\n".join(body))
        if text:
            segments.append(Segment("section", number, heading, text))
    return segments


# ---------- Images ----------

def _image(path: Path, ocr: bool, ocr_lang: str) -> list[Segment]:
    if not ocr:
        raise ExtractionError("Image files need OCR, but OCR is disabled (OCR_ENABLED=false).")
    from PIL import Image

    try:
        image = Image.open(path)
        image.load()
    except Exception as exc:
        raise ExtractionError("The image could not be opened.") from exc
    text = clean(_ocr_image(image.convert("RGB"), ocr_lang))
    return [Segment("image", 1, "", text)] if text else []
